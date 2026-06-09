# -*- coding: utf-8 -*-
"""
CALG multi-output coefficient, equation, and summary exporter.
"""

import pandas as pd
from pathlib import Path
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_squared_error, r2_score


SUMMARY_COLUMNS = [
    "Output Variable",
    "Type",
    "Within 0.01%",
    "Within 0.1%",
    "Within 1%",
    "Within 3%",
    "Within 5%",
    "Within 10%",
    "Beyond 10%",
    "Total Points",
]


def count_within(values, lower, upper):
    return sum(lower <= value <= upper for value in values)


def get_term_degree(column_name):
    term = str(column_name)
    if "*" not in term:
        return 1
    return term.count("*") + 1


def create_summary_row(output_column, model_type, y, y_pred):
    error_percent = ((y_pred - y) / y) * 100
    total_points = len(y)
    within_10 = count_within(error_percent, -10, 10)

    return {
        "Output Variable": output_column,
        "Type": model_type,
        "Within 0.01%": count_within(error_percent, -0.01, 0.01),
        "Within 0.1%": count_within(error_percent, -0.1, 0.1),
        "Within 1%": count_within(error_percent, -1, 1),
        "Within 3%": count_within(error_percent, -3, 3),
        "Within 5%": count_within(error_percent, -5, 5),
        "Within 10%": within_10,
        "Beyond 10%": total_points - within_10,
        "Total Points": total_points,
    }


def create_summary_rows(output_column, X, y):
    highest_degree = min(10, max(get_term_degree(column) for column in X.columns))
    summary_rows = []

    for degree in range(1, highest_degree + 1):
        degree_columns = [
            column
            for column in X.columns
            if get_term_degree(column) <= degree
        ]
        degree_model = LinearRegression()
        degree_model.fit(X[degree_columns], y)
        degree_pred = degree_model.predict(X[degree_columns])
        summary_rows.append(
            create_summary_row(output_column, f"Degree {degree}", y, degree_pred)
        )

    return summary_rows


def apply_summary_formatting(worksheet, summary_rows_count):
    header_fill = PatternFill("solid", fgColor="FFFFFF")
    model_fill = PatternFill("solid", fgColor="FFFF00")
    total_fill = PatternFill("solid", fgColor="F8CBAD")
    header_font = Font(bold=True, size=14)
    thin_side = Side(style="thin", color="000000")
    header_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center")

    total_points_column = len(SUMMARY_COLUMNS)
    for row in range(2, summary_rows_count + 2):
        for column in range(1, total_points_column + 1):
            cell = worksheet.cell(row=row, column=column)
            cell.fill = total_fill if column == total_points_column else model_fill
            cell.alignment = Alignment(horizontal="right")
        worksheet.cell(row=row, column=1).alignment = Alignment(horizontal="left")
        worksheet.cell(row=row, column=2).alignment = Alignment(horizontal="left")


def auto_adjust_column_widths(writer):
    for worksheet in writer.book.worksheets:
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            worksheet.column_dimensions[column_letter].width = max_length + 2


def sanitize_sheet_name(name):
    invalid_chars = ["\\", "/", "*", "[", "]", ":", "?"]
    sheet_name = str(name)
    for char in invalid_chars:
        sheet_name = sheet_name.replace(char, "_")
    return sheet_name[:31]


def get_unique_output_file(base_name):
    output_path = Path(base_name)
    if not output_path.exists():
        return str(output_path)

    index = 1
    while True:
        candidate = output_path.with_name(
            f"{output_path.stem}_{index}{output_path.suffix}"
        )
        if not candidate.exists():
            return str(candidate)
        index += 1


def build_equation(output_column, intercept, terms, coefficients):
    equation_lines = [f"{output_column} =", f"{intercept} +"]
    for index, (term, coefficient) in enumerate(zip(terms, coefficients)):
        suffix = " +" if index < len(terms) - 1 else ""
        equation_lines.append(f"{coefficient}*({term}){suffix}")
    return "\n".join(equation_lines)


def main():
    file = input("Enter Excel file path: ").strip().strip('"').strip("'")
    df = pd.read_excel(file)

    num_outputs = int(input("Enter number of output columns: "))
    if num_outputs < 1:
        raise ValueError("Number of output columns must be at least 1.")
    if num_outputs >= len(df.columns):
        raise ValueError("At least one input column must remain after output columns.")

    output_columns = list(df.columns[:num_outputs])
    input_columns = list(df.columns[num_outputs:])

    model_df = df[output_columns + input_columns].dropna().reset_index(drop=True)

    output_file = get_unique_output_file("CALG_Results.xlsx")
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for output_column in output_columns:
            model_df_output = model_df[model_df[output_column] != 0].reset_index(drop=True)
            y = model_df_output[output_column]
            X_output = model_df_output[input_columns]

            model = LinearRegression()
            model.fit(X_output, y)
            y_pred = model.predict(X_output)

            terms = list(X_output.columns)
            coefficients = list(model.coef_)
            intercept = model.intercept_

            summary_data = pd.DataFrame(
                create_summary_rows(output_column, X_output, y),
                columns=SUMMARY_COLUMNS,
            )
            metrics_data = pd.DataFrame(
                {
                    "Metric": [
                        "Output Variable",
                        "Number of Inputs",
                        "Number of Terms",
                        "R2 Score",
                        "MSE",
                        "Total Points",
                    ],
                    "Value": [
                        output_column,
                        len(input_columns),
                        len(terms),
                        r2_score(y, y_pred),
                        mean_squared_error(y, y_pred),
                        len(y),
                    ],
                }
            )

            coefficient_data = pd.DataFrame(
                {
                    "Variable": terms + ["Intercept"],
                    "Coefficient": coefficients + [intercept],
                }
            )

            equation = build_equation(output_column, intercept, terms, coefficients)
            equation_data = pd.DataFrame(
                {
                    "Variable": (
                        ["Output Variable", "Intercept"]
                        + terms
                        + ["Full Equation"]
                    ),
                    "Coefficient": (
                        [output_column, intercept]
                        + coefficients
                        + [equation]
                    ),
                }
            )

            sheet_prefix = sanitize_sheet_name(output_column)
            summary_sheet = sanitize_sheet_name(f"{sheet_prefix}_Summary")
            coefficient_sheet = sanitize_sheet_name(f"{sheet_prefix}_Coefficients")
            equation_sheet = sanitize_sheet_name(f"{sheet_prefix}_Equation")

            summary_data.to_excel(writer, sheet_name=summary_sheet, index=False)
            metrics_data.to_excel(
                writer,
                sheet_name=summary_sheet,
                startrow=len(summary_data) + 3,
                index=False,
            )
            coefficient_data.to_excel(
                writer,
                sheet_name=coefficient_sheet,
                index=False,
            )
            equation_data.to_excel(writer, sheet_name=equation_sheet, index=False)
            apply_summary_formatting(writer.book[summary_sheet], len(summary_data))

        auto_adjust_column_widths(writer)

    print(f"Results exported successfully to {output_file}")


if __name__ == "__main__":
    main()