from pathlib import Path
from uuid import uuid4

import numpy as np
import pandas as pd
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_squared_error, r2_score
from werkzeug.utils import secure_filename


BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"
ALLOWED_EXTENSIONS = {".xlsx", ".xls"}

SUMMARY_COLUMNS = [
    "Output Variable",
    "Type",
    "Within 0.01%",
    "Within 0.1%",
    "Within 1%",
    "Within 3%",
    "Within 5%",
    "Within 10%",
    "Beyond 10%",
    "Total Points",
]

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 250 * 1024 * 1024


def count_within(values, lower, upper):
    return sum(lower <= value <= upper for value in values)


def get_term_degree(column_name):
    term = str(column_name)
    if "*" not in term:
        return 1
    return term.count("*") + 1


def create_summary_row(output_column, model_type, y, y_pred):
    error_percent = ((y_pred - y) / y) * 100
    total_points = len(y)
    within_10 = count_within(error_percent, -10, 10)

    return {
        "Output Variable": output_column,
        "Type": model_type,
        "Within 0.01%": count_within(error_percent, -0.01, 0.01),
        "Within 0.1%": count_within(error_percent, -0.1, 0.1),
        "Within 1%": count_within(error_percent, -1, 1),
        "Within 3%": count_within(error_percent, -3, 3),
        "Within 5%": count_within(error_percent, -5, 5),
        "Within 10%": within_10,
        "Beyond 10%": total_points - within_10,
        "Total Points": total_points,
    }


def create_summary_rows(output_column, X, y):
    highest_degree = min(10, max(get_term_degree(column) for column in X.columns))

    # Produce summary only for the highest detected degree (not incremental degrees)
    degree = highest_degree
    degree_columns = [
        column
        for column in X.columns
        if get_term_degree(column) <= degree
    ]
    degree_model = LinearRegression()
    degree_model.fit(X[degree_columns], y)
    degree_pred = degree_model.predict(X[degree_columns])

    return [create_summary_row(output_column, f"Degree {degree}", y, degree_pred)]


def apply_summary_formatting(worksheet, summary_rows_count):
    header_fill = PatternFill("solid", fgColor="FFFFFF")
    model_fill = PatternFill("solid", fgColor="FFFF00")
    total_fill = PatternFill("solid", fgColor="F8CBAD")
    header_font = Font(bold=True, size=14)
    thin_side = Side(style="thin", color="000000")
    header_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = header_border
        cell.alignment = Alignment(horizontal="center")

    total_points_column = len(SUMMARY_COLUMNS)
    for row in range(2, summary_rows_count + 2):
        for column in range(1, total_points_column + 1):
            cell = worksheet.cell(row=row, column=column)
            cell.fill = total_fill if column == total_points_column else model_fill
            cell.alignment = Alignment(horizontal="right")
        worksheet.cell(row=row, column=1).alignment = Alignment(horizontal="left")
        worksheet.cell(row=row, column=2).alignment = Alignment(horizontal="left")


def auto_adjust_column_widths(writer):
    for worksheet in writer.book.worksheets:
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            worksheet.column_dimensions[column_letter].width = max_length + 2


def sanitize_sheet_name(name):
    invalid_chars = ["\\", "/", "*", "[", "]", ":", "?"]
    sheet_name = str(name)
    for char in invalid_chars:
        sheet_name = sheet_name.replace(char, "_")
    return sheet_name[:31]


def build_equation(output_column, intercept, terms, coefficients):
    equation_lines = [f"{output_column} =", f"{intercept} +"]
    for index, (term, coefficient) in enumerate(zip(terms, coefficients)):
        suffix = " +" if index < len(terms) - 1 else ""
        equation_lines.append(f"{coefficient}*({term}){suffix}")
    return "\n".join(equation_lines)


def error_response(message, status_code=400):
    return jsonify({"error": message}), status_code


def allowed_excel(filename):
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def generate_workbook(input_file, num_outputs, output_file):
    df = pd.read_excel(input_file)
    if df.empty:
        raise ValueError("Uploaded Excel file has no data.")
    if num_outputs < 1:
        raise ValueError("Number of output columns must be at least 1.")
    if num_outputs >= len(df.columns):
        raise ValueError("Number of output columns must be less than total columns.")

    output_columns = list(df.columns[:num_outputs])
    input_columns = list(df.columns[num_outputs:])
    model_df = df[output_columns + input_columns].dropna().reset_index(drop=True)

    if model_df.empty:
        raise ValueError("No usable rows remain after removing blank values.")

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for output_column in output_columns:
            model_df_output = model_df[model_df[output_column] != 0].reset_index(drop=True)
            if model_df_output.empty:
                raise ValueError(
                    f"Output column '{output_column}' has no non-zero rows to fit."
                )

            y = model_df_output[output_column]
            X_output = model_df_output[input_columns].apply(pd.to_numeric, errors="coerce")
            valid_rows = X_output.notna().all(axis=1) & pd.to_numeric(
                y, errors="coerce"
            ).notna()
            X_output = X_output.loc[valid_rows].reset_index(drop=True)
            y = pd.to_numeric(y.loc[valid_rows], errors="coerce").reset_index(drop=True)

            if X_output.empty or len(y) == 0:
                raise ValueError(
                    f"Output column '{output_column}' has no numeric rows to fit."
                )
            if not np.isfinite(X_output.to_numpy()).all() or not np.isfinite(y.to_numpy()).all():
                raise ValueError(
                    f"Output column '{output_column}' contains invalid numeric values."
                )

            try:
                model = LinearRegression()
                model.fit(X_output, y)
                y_pred = model.predict(X_output)
                summary_rows = create_summary_rows(output_column, X_output, y)
            except Exception as exc:
                raise RuntimeError(
                    f"Regression failure for output column '{output_column}': {exc}"
                ) from exc

            terms = list(X_output.columns)
            coefficients = list(model.coef_)
            intercept = model.intercept_

            summary_data = pd.DataFrame(summary_rows, columns=SUMMARY_COLUMNS)
            metrics_data = pd.DataFrame(
                {
                    "Metric": [
                        "Output Variable",
                        "Number of Inputs",
                        "Number of Terms",
                        "R2 Score",
                        "MSE",
                        "Total Points",
                    ],
                    "Value": [
                        output_column,
                        len(input_columns),
                        len(terms),
                        r2_score(y, y_pred),
                        mean_squared_error(y, y_pred),
                        len(y),
                    ],
                }
            )

            coefficient_data = pd.DataFrame(
                {
                    "Variable": terms + ["Intercept"],
                    "Coefficient": coefficients + [intercept],
                }
            )

            equation = build_equation(output_column, intercept, terms, coefficients)
            equation_data = pd.DataFrame(
                {
                    "Variable": (
                        ["Output Variable", "Intercept"] + terms + ["Full Equation"]
                    ),
                    "Coefficient": (
                        [output_column, intercept] + coefficients + [equation]
                    ),
                }
            )

            sheet_prefix = sanitize_sheet_name(output_column)
            summary_sheet = sanitize_sheet_name(f"{sheet_prefix}_Summary")
            coefficient_sheet = sanitize_sheet_name(f"{sheet_prefix}_Coefficients")
            equation_sheet = sanitize_sheet_name(f"{sheet_prefix}_Equation")

            summary_data.to_excel(writer, sheet_name=summary_sheet, index=False)
            metrics_data.to_excel(
                writer,
                sheet_name=summary_sheet,
                startrow=len(summary_data) + 3,
                index=False,
            )
            coefficient_data.to_excel(
                writer,
                sheet_name=coefficient_sheet,
                index=False,
            )
            equation_data.to_excel(writer, sheet_name=equation_sheet, index=False)
            apply_summary_formatting(writer.book[summary_sheet], len(summary_data))

        auto_adjust_column_widths(writer)


@app.route("/")
def index():
    return render_template("index.html")


@app.post("/generate")
def generate():
    upload = request.files.get("file")
    if upload is None or upload.filename == "":
        return error_response("No file selected.")
    if not allowed_excel(upload.filename):
        return error_response("Invalid Excel file. Upload a .xlsx or .xls file.")

    try:
        num_outputs = int(request.form.get("num_outputs", ""))
    except ValueError:
        return error_response("Invalid number of outputs.")

    if num_outputs < 1:
        return error_response("Invalid number of outputs.")

    UPLOAD_DIR.mkdir(exist_ok=True)
    OUTPUT_DIR.mkdir(exist_ok=True)

    safe_name = secure_filename(upload.filename) or "upload.xlsx"
    request_id = uuid4().hex
    input_file = UPLOAD_DIR / f"{request_id}_{safe_name}"
    output_file = OUTPUT_DIR / f"{request_id}_CALG_Results.xlsx"

    try:
        upload.save(input_file)
        generate_workbook(input_file, num_outputs, output_file)
    except ValueError as exc:
        return error_response(str(exc))
    except RuntimeError as exc:
        return error_response(str(exc), 500)
    except Exception as exc:
        return error_response(f"Workbook generation failure: {exc}", 500)
    finally:
        input_file.unlink(missing_ok=True)

    return send_file(
        output_file,
        as_attachment=True,
        download_name="CALG_Results.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000)
